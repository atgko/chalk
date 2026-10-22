"""Schedules with no dates (UAT of PR #18: IS 6490's "Course Outline" is
"Meeting 1" to "Meeting 10" with no dates): reading them from Word and
PDF, dating them, and rolling them over."""

import datetime as dt

from docx import Document

from chalk.extractors._docx_schedule import is_bare_week_number, parse_week_label
from chalk.extractors.docx_extractor import extract_course_data
from chalk.extractors.undated_schedule import (
    date_undated_schedule,
    default_week1,
    provisional_week1,
    syllabus_lists_breaks,
)
from chalk.models import UniversityDate
from chalk.pdf_import.tables import assemble, freeze
from chalk.rollover.docx_rollover import write_rolled_over_docx
from tests.fixtures.docx_builder import _add_front_matter

_CALENDARS = {
    "terms": [
        {
            "term": "Fall 2026",
            "start": "2026-08-24",
            "end": "2026-12-10",
            "no_class_dates": [
                {"label": "Labor Day", "date": "2026-09-07"},
                {"label": "Fall Break", "date_start": "2026-10-10", "date_end": "2026-10-18"},
            ],
        }
    ]
}


def _outline_syllabus(path, *, header="Meeting", meetings=10, meeting_pattern="Online async"):
    doc = Document()
    _add_front_matter(doc, meeting_pattern=meeting_pattern)
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, (header, "Managerial Decision Focus", "AI Capability Focus")):
        cell.text = text
    for number in range(1, meetings + 1):
        cells = table.add_row().cells
        cells[0].text = str(number)
        cells[1].text = f"Decision {number}"
        cells[2].text = f"Capability {number}"
    doc.save(str(path))
    return path


def _regular(course_data):
    return [week for week in course_data.weeks if not week.is_break]


# ---- Labels -------------------------------------------------------------------------


def test_undated_labels_are_read_only_when_asked():
    assert parse_week_label("10") is None
    label = parse_week_label("Meeting 3", undated=True)
    assert (label.week_number, label.is_dated, label.prefix) == (3, False, "Meeting ")
    assert parse_week_label("Week 2 (8/31)").is_dated


def test_a_bare_week_number_is_the_whole_cell():
    assert is_bare_week_number("10") and is_bare_week_number(" Meeting 4 ")
    assert not is_bare_week_number("1 Exam (100 points each)")
    assert not is_bare_week_number("Week 1 (8/24)")


# ---- Word ---------------------------------------------------------------------------


def test_a_meeting_outline_is_read_as_the_schedule(tmp_path):
    course_data = extract_course_data(_outline_syllabus(tmp_path / "outline.docx"))

    weeks = _regular(course_data)
    assert course_data.course.undated_schedule
    assert [w.week_number for w in weeks] == list(range(1, 11))
    assert weeks[0].topics == ["Decision 1", "Capability 1"]
    assert weeks[1].date - weeks[0].date == dt.timedelta(weeks=1)  # provisional, a week apart


def test_a_dated_schedule_is_not_marked_undated(tmp_path):
    from tests.fixtures import docx_builder

    course_data = extract_course_data(docx_builder.build_minimal_syllabus(tmp_path / "s.docx"))
    assert not course_data.course.undated_schedule


def test_a_numbered_table_without_a_schedule_header_is_not_a_schedule(tmp_path):
    from chalk.errors import AmbiguousTableError

    path = _outline_syllabus(tmp_path / "list.docx", header="Policy")
    try:
        extract_course_data(path)
    except AmbiguousTableError as exc:
        assert exc.user_message.startswith("No schedule table found.")
    else:  # pragma: no cover - the assertion is the point
        raise AssertionError("a numbered policy list was read as a schedule")


# ---- Dating -------------------------------------------------------------------------


def test_provisional_week1_is_a_monday_in_the_terms_season():
    assert provisional_week1("Fall 2026", 2026) == dt.date(2026, 8, 24)
    assert provisional_week1("Spring 2027", 2027) == dt.date(2027, 1, 11)
    assert provisional_week1("Summer 2027", 2027).month == 5


def test_default_week1_is_the_terms_first_class_day(tmp_path):
    course = extract_course_data(_outline_syllabus(tmp_path / "o.docx", meeting_pattern="Wednesdays 6-9 pm")).course
    assert default_week1(course, _CALENDARS) == dt.date(2026, 8, 26)  # first Wednesday from Aug 24
    # Not in the calendar: the season's first Monday, then the first Wednesday after it.
    assert default_week1(course.model_copy(update={"term": "Fall 2031"}), _CALENDARS) == dt.date(2031, 8, 27)


def test_dating_skips_a_multi_day_break_and_adds_its_row(tmp_path):
    course_data = extract_course_data(_outline_syllabus(tmp_path / "o.docx"))

    dated = date_undated_schedule(course_data, dt.date(2026, 8, 24), _CALENDARS, skip_breaks=True)

    dates = [w.date for w in _regular(dated)]
    assert dates[6] == dt.date(2026, 10, 5)
    assert dates[7] == dt.date(2026, 10, 19)  # Oct 12 is in Fall Break
    assert dates[1] == dt.date(2026, 8, 31) and dates[2] == dt.date(2026, 9, 7)  # Labor Day doesn't move it
    [fall_break] = [w for w in dated.weeks if w.is_break]
    assert fall_break.label == "Fall Break (10/10 - 10/18)"
    assert dated.weeks.index(fall_break) == 7
    assert (dated.course.term_start, dated.course.term_end) == (dates[0], dates[-1])


def test_changing_week1_redates_and_doesnt_duplicate_the_break(tmp_path):
    course_data = extract_course_data(_outline_syllabus(tmp_path / "o.docx"))
    first = date_undated_schedule(course_data, dt.date(2026, 8, 24), _CALENDARS, skip_breaks=True)

    again = date_undated_schedule(first, dt.date(2026, 8, 31), _CALENDARS, skip_breaks=True)

    assert _regular(again)[0].date == dt.date(2026, 8, 31)
    assert sum(1 for w in again.weeks if w.is_break) == 1


def test_without_breaks_the_class_meets_through_them(tmp_path):
    course_data = extract_course_data(_outline_syllabus(tmp_path / "o.docx"))
    with_break = date_undated_schedule(course_data, dt.date(2026, 8, 24), _CALENDARS, skip_breaks=True)

    dated = date_undated_schedule(with_break, dt.date(2026, 9, 28), _CALENDARS, skip_breaks=False)

    assert not any(w.is_break for w in dated.weeks)  # the earlier break row is gone too
    dates = [w.date for w in dated.weeks]
    assert dates[2] == dt.date(2026, 10, 12)  # in Fall Break, which this course doesn't take
    assert dates == [dt.date(2026, 9, 28) + dt.timedelta(weeks=n) for n in range(10)]


def test_a_syllabus_lists_breaks_only_when_it_mentions_one(tmp_path):
    course_data = extract_course_data(_outline_syllabus(tmp_path / "o.docx"))
    assert not syllabus_lists_breaks(course_data)

    def with_dates(*dates):
        return course_data.model_copy(update={"university_dates": list(dates)})

    assert not syllabus_lists_breaks(with_dates(UniversityDate(event="Last day to drop", date=dt.date(2026, 9, 4))))
    assert syllabus_lists_breaks(with_dates(UniversityDate(event="Fall break", date=dt.date(2026, 10, 12))))
    span = UniversityDate(event="No classes", date_start=dt.date(2026, 10, 10), date_end=dt.date(2026, 10, 18))
    assert syllabus_lists_breaks(with_dates(span))


# ---- Rollover writer --------------------------------------------------------------


def test_rolling_over_an_outline_keeps_its_numbering(tmp_path):
    source = _outline_syllabus(tmp_path / "o.docx", meetings=3)
    course_data = extract_course_data(source)
    weeks = _regular(course_data)
    extra = weeks[-1].model_copy(update={"week_number": 4, "date": weeks[-1].date + dt.timedelta(weeks=1), "topics": ["New"]})
    rolled = course_data.model_copy(update={"weeks": [*weeks, extra]})

    write_rolled_over_docx(source, rolled, tmp_path / "out.docx")

    labels = [row.cells[0].text for row in Document(str(tmp_path / "out.docx")).tables[0].rows]
    assert labels == ["Meeting", "1", "2", "3", "4"]


# ---- PDF ----------------------------------------------------------------------------


def test_a_numbered_row_continues_the_schedule_on_the_next_page():
    """IS 6490: page 10 holds the end of Meeting 9 and all of Meeting 10."""
    header = [["Meeting"], ["Focus"], ["Outcome"]]
    page_one = freeze([header, [["9"], ["Apply the complete managerial"], ["Demonstrate"]]])
    page_two = freeze(
        [[[], ["opportunities across business scenarios."], []], [["10"], ["Executive Presentation"], ["Defend"]]],
        is_real=False,
    )

    [table] = assemble([page_one, page_two])

    assert [row[0] for row in table.rows] == [("Meeting",), ("9",), ("10",)]
    assert table.rows[1][1] == ("Apply the complete managerial opportunities across business scenarios.",)
