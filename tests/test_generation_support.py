"""Tests for the generation building blocks: prompt templates, source
context, cost calculation, and slide normalization."""

import pytest
from docx import Document

from chalk.costs import cost_usd, estimate_max_cost, estimate_tokens, rate_for
from chalk.errors import GenerationError, PromptTemplateError
from chalk.generation.prompts import fill_template, load_template, template_path
from chalk.generation.slides import normalize_slides
from chalk.generation.source_context import (
    NO_SOURCES_TEXT,
    build_source_context,
    read_source_text,
)
from chalk.generation.specs import SPECS
from chalk.project import RESOURCES_DIR
from tests.fixtures.pdf_builder import build_text_pdf

CONFIG = {
    "cost_rates": {
        "openai": {"gpt-4o": {"input_per_1k": 0.0025, "output_per_1k": 0.010}},
        "anthropic": {"claude-sonnet-5": {"input_per_1k": 0.003, "output_per_1k": 0.015}},
        "local": {"default": {"input_per_1k": 0.0, "output_per_1k": 0.0}},
    }
}

# ---- Prompt templates ------------------------------------------------------------------


@pytest.mark.parametrize("key", list(SPECS))
def test_every_bundled_template_has_its_required_variables(key, tmp_path):
    template = load_template(tmp_path, SPECS[key])  # tmp_path has no copy -> bundled default
    assert "{course_title}" in template


def test_a_project_copy_takes_precedence_over_the_bundled_template(tmp_path):
    spec = SPECS["summary"]
    custom = "Mine: {course_title} {learning_objectives} {source_material_context} {week_number} {topics}"
    (tmp_path / "summary.txt").write_text(custom, encoding="utf-8")

    assert template_path(tmp_path, spec) == tmp_path / "summary.txt"
    assert load_template(tmp_path, spec) == custom


def test_a_template_missing_a_variable_names_it_and_the_file(tmp_path):
    (tmp_path / "quiz.txt").write_text("Quiz for {course_title} on {topics}", encoding="utf-8")

    with pytest.raises(PromptTemplateError) as exc_info:
        load_template(tmp_path, SPECS["quiz"])

    message = exc_info.value.user_message
    assert str(tmp_path / "quiz.txt") in message
    assert "{week_number}" in message and "{question_count}" in message
    assert "{course_title}" not in message


def test_an_unreadable_template_is_a_plain_error(tmp_path):
    (tmp_path / "quiz.txt").write_bytes(b"\xff\xfe\x00bad")
    with pytest.raises(PromptTemplateError, match="couldn't be read"):
        load_template(tmp_path, SPECS["quiz"])


def test_a_template_missing_everywhere_is_a_plain_error(tmp_path):
    with pytest.raises(PromptTemplateError):
        load_template(tmp_path, SPECS["quiz"], resources_dir=tmp_path / "nowhere")


def test_fill_template_leaves_unknown_braces_alone():
    filled = fill_template(":::: {.columns} {course_title} {unknown} {", {"course_title": "IS 6640"})
    assert filled == ":::: {.columns} IS 6640 {unknown} {"


# ---- Source context -------------------------------------------------------------------


def test_no_selected_sources_uses_the_fallback_text(tmp_project):
    assert build_source_context(tmp_project, []) == NO_SOURCES_TEXT


def test_reads_txt_md_docx_and_pdf(tmp_project):
    source = tmp_project.source_dir
    (source / "a.txt").write_text("plain notes", encoding="utf-8")
    (source / "b.md").write_text("# Heading\nmarkdown notes", encoding="utf-8")
    document = Document()
    document.add_paragraph("docx paragraph")
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "cell text"
    document.save(str(source / "c.docx"))
    build_text_pdf(source / "d.pdf", "Subnetting (CIDR) basics")

    context = build_source_context(tmp_project, ["a.txt", "b.md", "c.docx", "d.pdf"])

    for expected in ("--- Source: a.txt ---", "plain notes", "markdown notes", "docx paragraph",
                     "cell text", "Subnetting (CIDR) basics"):
        assert expected in context


def test_pdf_without_text_or_corrupt_pdf_yields_empty_text(tmp_path):
    corrupt = tmp_path / "bad.pdf"
    corrupt.write_bytes(b"%PDF-1.4 not really")
    assert read_source_text(corrupt) == ""


def test_file_with_no_extractable_text_says_so(tmp_project):
    (tmp_project.source_dir / "scan.pdf").write_bytes(b"%PDF-1.4 garbage")
    assert "(No extractable text in this file.)" in build_source_context(tmp_project, ["scan.pdf"])


def test_unknown_source_names_are_rejected_not_resolved_as_paths(tmp_project):
    with pytest.raises(GenerationError, match=r"\.\./secrets\.txt"):
        build_source_context(tmp_project, ["../secrets.txt"])


def test_context_is_truncated_to_the_token_budget(tmp_project):
    (tmp_project.source_dir / "long.txt").write_text("x" * 5000, encoding="utf-8")

    context = build_source_context(tmp_project, ["long.txt"], token_budget=100)

    assert context.startswith("Source materials supplied by the instructor:")
    assert len(context.split("\n\n[Source material truncated")[0]) == 400
    assert "more characters omitted" in context


# ---- Costs ------------------------------------------------------------------------------


def test_rates_are_looked_up_per_provider_and_model():
    assert rate_for(CONFIG, {"LLM_PROVIDER": "openai", "LLM_MODEL": "gpt-4o"})["output_per_1k"] == 0.010
    assert rate_for(CONFIG, {"LLM_PROVIDER": "anthropic", "LLM_MODEL": "claude-sonnet-5"})["input_per_1k"] == 0.003


def test_local_endpoints_always_use_the_local_default_rate():
    env = {"LLM_PROVIDER": "openai", "LLM_BASE_URL": "http://localhost:11434/v1", "LLM_MODEL": "gpt-4o"}
    assert rate_for(CONFIG, env) == {"input_per_1k": 0.0, "output_per_1k": 0.0}
    assert rate_for({}, env) == {"input_per_1k": 0.0, "output_per_1k": 0.0}


def test_hosted_model_without_a_rate_is_unknown_not_free():
    rate = rate_for(CONFIG, {"LLM_PROVIDER": "openai", "LLM_MODEL": "gpt-typo"})
    assert rate is None
    assert cost_usd(rate, 1000, 1000) is None


def test_cost_and_estimate_arithmetic():
    rate = CONFIG["cost_rates"]["openai"]["gpt-4o"]
    assert cost_usd(rate, 2000, 500) == pytest.approx(0.005 + 0.005)
    assert estimate_tokens("a" * 400) == 100
    assert estimate_tokens("") == 1
    assert estimate_max_cost(rate, "a" * 4000, 2000) == pytest.approx(0.0025 + 0.02)


# ---- Slides -------------------------------------------------------------------------------


def _normalize(text):
    return normalize_slides(text, default_title="Week 3: Subnetting", author="Dave Norwood", banner="BANNER")


def test_slides_get_yaml_first_numbered_slides_and_banner_comment():
    deck = _normalize("## Intro\n- one\n\n## Details\n- two")
    lines = deck.splitlines()
    assert lines[:4] == ["---", 'title: "Week 3: Subnetting"', 'author: "Dave Norwood"', "---"]
    assert "<!-- BANNER -->" in lines
    assert deck.index("<!-- Slide 1 -->\n## Intro") < deck.index("<!-- Slide 2 -->\n## Details")


def test_slides_drop_subtitle_and_keep_existing_title_and_author():
    deck = _normalize('---\ntitle: "Mine"\nsubtitle: "nope"\nauthor: "Me"\n---\n## A\n- x')
    assert 'title: "Mine"' in deck and 'author: "Me"' in deck
    assert "subtitle" not in deck
    assert "Week 3: Subnetting" not in deck


def test_slides_convert_bare_h1_and_renumber_stale_slide_comments():
    deck = _normalize("<!-- Slide 7 -->\n# Title slide\n#No space\n## Next\n###Keep h3")
    assert "\n# " not in deck
    assert "## Title slide" in deck and "## No space" in deck
    assert [line for line in deck.splitlines() if line.startswith("<!-- Slide")] == [
        "<!-- Slide 1 -->", "<!-- Slide 2 -->", "<!-- Slide 3 -->",
    ]
    assert "###Keep h3" in deck


def test_slides_with_unterminated_front_matter_treat_it_as_body():
    deck = _normalize("---\ntitle: never closed\n## A")
    assert deck.startswith('---\ntitle: "Week 3: Subnetting"')


def test_bundled_resources_dir_holds_all_prompt_templates():
    assert {p.stem for p in (RESOURCES_DIR / "prompts").glob("*.txt")} == set(SPECS)
