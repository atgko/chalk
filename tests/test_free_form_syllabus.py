"""Real syllabi rarely use Chalk's "Label: value" header convention
(teammate feedback, Sept 2026). These cover the looser extraction end to
end, including rolling such a syllabus over — where the term lives in the
title line rather than a "Term:" line."""

from docx import Document

from chalk.extractors import docx_extractor, markdown_extractor
from chalk.rollover.docx_rollover import write_rolled_over_docx
from chalk.rollover.markdown_rollover import write_rolled_over_markdown
from chalk.rollover.preview import roll_over_course
from tests.fixtures.docx_builder import (
    build_syllabus_with_course_info_table,
    build_syllabus_with_free_form_header,
)
from tests.test_rollover_season_change import CALENDARS


def test_free_form_header_extracts_course_info_and_schedule(tmp_path):
    path = build_syllabus_with_free_form_header(tmp_path / "syllabus.docx")

    course_data = docx_extractor.extract_course_data(path)

    course = course_data.course
    assert course.title == "Networking and Servers – Online"
    assert course.number == "IS 6640"
    assert course.term == "Fall 2026"
    assert course.instructor == "Jane Example"
    assert course.duration_weeks == 3
    assert [w.topics for w in course_data.weeks] == [
        ["Network+ Mod 1"],
        ["Network+ Mod 2"],
        ["Network+ Mod 3"],
    ]


def test_course_info_table_is_read(tmp_path):
    path = build_syllabus_with_course_info_table(tmp_path / "syllabus.docx")

    course = docx_extractor.extract_course_data(path).course

    assert (course.title, course.number, course.term, course.credits, course.instructor) == (
        "Networking and Servers",
        "IS 6640",
        "Fall 2026",
        3,
        "Jane Example",
    )


def test_docx_rollover_updates_the_term_inside_the_title_line(tmp_path):
    source = build_syllabus_with_free_form_header(tmp_path / "syllabus.docx")
    new_course_data, _ = roll_over_course(
        docx_extractor.extract_course_data(source), target_term="Spring 2027", calendars=CALENDARS
    )
    output = tmp_path / "out" / "syllabus.docx"

    write_rolled_over_docx(source, new_course_data, output)

    paragraphs = [p.text for p in Document(str(output)).paragraphs]
    assert "Networking and Servers – Online Spring 2027" in paragraphs
    reextracted = docx_extractor.extract_course_data(output)
    assert reextracted.course.term == "Spring 2027"
    assert reextracted.weeks[0].date.isoformat() == "2027-01-11"


def test_markdown_rollover_updates_the_term_inside_the_heading(tmp_path):
    source = tmp_path / "syllabus.md"
    source.write_text(
        "# IS 4490: Emerging Technologies — Fall 2026\n\n"
        "| Week | Dates | Topic | Major Work |\n| --- | --- | --- | --- |\n"
        "| 1 | 8/24 - 8/30 | Course Introduction |  |\n",
        encoding="utf-8",
    )
    new_course_data, _ = roll_over_course(
        markdown_extractor.extract_course_data(source),
        target_term="Spring 2027",
        calendars=CALENDARS,
    )
    output = tmp_path / "out" / "syllabus.md"

    write_rolled_over_markdown(source, new_course_data, output)

    assert output.read_text(encoding="utf-8").startswith(
        "# IS 4490: Emerging Technologies — Spring 2027\n"
    )
    assert markdown_extractor.extract_course_data(output).course.term == "Spring 2027"


def test_labeled_term_line_keeps_its_own_label_on_rollover(tmp_path):
    source = tmp_path / "syllabus.md"
    source.write_text(
        "Title: Emerging Technologies\nSemester: Fall 2026\n\n"
        "| Week | Dates | Topic | Major Work |\n| --- | --- | --- | --- |\n"
        "| 1 | 8/24 - 8/30 | Course Introduction |  |\n",
        encoding="utf-8",
    )
    new_course_data, _ = roll_over_course(
        markdown_extractor.extract_course_data(source),
        target_term="Spring 2027",
        calendars=CALENDARS,
    )
    output = tmp_path / "out" / "syllabus.md"

    write_rolled_over_markdown(source, new_course_data, output)

    assert "Semester: Spring 2027" in output.read_text(encoding="utf-8").splitlines()


def test_docx_term_split_across_runs_is_still_rolled_over(tmp_path):
    # Word often splits a line into several runs (spell-check, edits).
    source = build_syllabus_with_free_form_header(tmp_path / "syllabus.docx")
    document = Document(str(source))
    title = next(p for p in document.paragraphs if "Fall 2026" in p.text)
    title.runs[0].text = "Networking and Servers – Online Fall "
    title.add_run("2026").bold = True
    document.save(str(source))
    new_course_data, _ = roll_over_course(
        docx_extractor.extract_course_data(source), target_term="Spring 2027", calendars=CALENDARS
    )
    output = tmp_path / "out" / "syllabus.docx"

    write_rolled_over_docx(source, new_course_data, output)

    assert docx_extractor.extract_course_data(output).course.term == "Spring 2027"
