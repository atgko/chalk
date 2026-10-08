"""End-to-end season-change rollovers through the Word and markdown
writers: extract -> roll_over_course -> write -> read the result back."""

from docx import Document

from chalk.extractors import docx_extractor, markdown_extractor
from chalk.rollover.docx_rollover import write_rolled_over_docx
from chalk.rollover.markdown_rollover import write_rolled_over_markdown
from chalk.rollover.preview import roll_over_course
from tests.fixtures.docx_builder import build_syllabus_with_full_schedule
from tests.fixtures.md_builder import build_minimal_syllabus as build_minimal_md_syllabus
from tests.test_rollover_season_change import CALENDARS


def _roll_docx(tmp_path, target_term: str, **builder_kwargs):
    source = build_syllabus_with_full_schedule(tmp_path / "syllabus.docx", **builder_kwargs)
    new_course_data, _ = roll_over_course(
        docx_extractor.extract_course_data(source), target_term=target_term, calendars=CALENDARS
    )
    output = tmp_path / "out" / "syllabus.docx"
    write_rolled_over_docx(source, new_course_data, output)
    return output


def _schedule_labels(docx_path) -> list[str]:
    table = Document(str(docx_path)).tables[0]
    return [row.cells[0].text for row in table.rows[1:]]


def test_docx_fall_to_spring_relabels_and_moves_the_break_row(tmp_path):
    output = _roll_docx(tmp_path, "Spring 2027")

    labels = _schedule_labels(output)
    assert not any("Fall Break" in label for label in labels)
    assert labels.index("Week 8 (3/1)") + 1 == labels.index("Spring Break (3/6 - 3/14)")
    assert labels.index("Spring Break (3/6 - 3/14)") + 1 == labels.index("Week 9 (3/8)")


def test_docx_break_row_keeps_its_other_cell_content(tmp_path):
    output = _roll_docx(tmp_path, "Spring 2027")

    table = Document(str(output)).tables[0]
    row = next(r for r in table.rows if r.cells[0].text.startswith("Spring Break"))
    assert row.cells[1].text == "No class"


def test_docx_fall_to_summer_removes_the_break_row(tmp_path):
    output = _roll_docx(tmp_path, "Summer 2027")

    labels = _schedule_labels(output)
    assert len(labels) == 10
    assert all(label.startswith("Week") for label in labels)


def test_docx_spring_to_fall_adds_a_row_for_each_new_break(tmp_path):
    output = _roll_docx(
        tmp_path,
        "Fall 2027",
        term="Spring 2027",
        week1=(2027, 1, 11),
        weeks=16,
        break_row=("Spring Break (3/6 - 3/14)", "No class"),
        break_after_week=8,
    )

    labels = _schedule_labels(output)
    assert labels.index("Week 7 (10/4)") + 1 == labels.index("Fall Break (10/9 - 10/17)")
    assert labels.index("Week 14 (11/22)") + 1 == labels.index("Thanksgiving Break (11/25 - 11/28)")
    assert len(labels) == 18


def test_docx_rolled_over_file_re_extracts_with_the_new_break(tmp_path):
    output = _roll_docx(tmp_path, "Spring 2027")

    reextracted = docx_extractor.extract_course_data(output)
    assert reextracted.course.term == "Spring 2027"
    assert [w.label for w in reextracted.weeks if w.is_break] == ["Spring Break (3/6 - 3/14)"]


def test_docx_same_season_rollover_still_keeps_the_break_in_place(tmp_path):
    output = _roll_docx(tmp_path, "Fall 2027")

    labels = _schedule_labels(output)
    assert labels.index("Week 7 (10/4)") + 1 == labels.index("Fall Break (10/9 - 10/17)")


def test_docx_university_dates_row_is_renamed_for_the_new_season(tmp_path):
    output = _roll_docx(
        tmp_path,
        "Spring 2027",
        university_dates=[
            ("Classes begin", "8/24/2026"),
            ("Fall Break", "10/10 - 10/18"),
            ("Classes end", "12/10/2026"),
        ],
    )

    dates_table = Document(str(output)).tables[1]
    rows = [(row.cells[0].text, row.cells[1].text) for row in dates_table.rows[1:]]
    assert rows == [
        ("Classes begin", "1/11/2027"),
        ("Spring Break", "3/6 - 3/14"),
        ("Classes end", "4/27/2027"),
    ]



def test_docx_holiday_row_is_renamed_for_the_new_season(tmp_path):
    output = _roll_docx(
        tmp_path,
        "Spring 2027",
        university_dates=[
            ("Classes begin", "8/24/2026"),
            ("Labor Day", "9/7/2026"),
            ("Fall Break", "10/10 - 10/18"),
            ("Classes end", "12/10/2026"),
        ],
    )

    dates_table = Document(str(output)).tables[1]
    rows = [(row.cells[0].text, row.cells[1].text) for row in dates_table.rows[1:]]
    assert rows[1] == ("Martin Luther King Jr. Day", "1/18/2027")

def test_markdown_fall_to_spring_renames_breaks_in_schedule_and_dates(tmp_path):
    source = build_minimal_md_syllabus(tmp_path / "syllabus.md")
    course_data = markdown_extractor.extract_course_data(source)
    # The minimal fixture's two weeks end in January — stretch the course so
    # Spring Break falls inside it.
    new_course_data, _ = roll_over_course(
        course_data,
        target_term="Spring 2027",
        calendars=CALENDARS,
        target_duration_weeks=10,
        llm_generate_topics=False,
    )
    output = tmp_path / "out" / "syllabus.md"
    write_rolled_over_markdown(source, new_course_data, output)

    text = output.read_text(encoding="utf-8")
    assert "Fall Break" not in text
    assert "| - | 3/6 - 3/14 | Spring Break |  |" in text
    assert "| Spring Break | 3/6 - 3/14 |" in text


def test_docx_rollover_with_a_longer_course_writes_and_rereads_every_week(tmp_path):
    source = build_syllabus_with_full_schedule(tmp_path / "syllabus.docx")
    course_data = docx_extractor.extract_course_data(source)
    weeks_before = sum(1 for week in course_data.weeks if not week.is_break)
    new_course_data, _ = roll_over_course(
        course_data,
        target_term="Fall 2027",
        calendars=CALENDARS,
        target_duration_weeks=weeks_before + 2,
        llm_generate_topics=False,
    )
    output = tmp_path / "out" / "syllabus.docx"
    write_rolled_over_docx(source, new_course_data, output)

    reread = [w for w in docx_extractor.extract_course_data(output).weeks if not w.is_break]
    expected = {w.week_number: w.date for w in new_course_data.weeks if not w.is_break}
    assert {w.week_number: w.date for w in reread} == expected
    assert reread[-1].notes == "No LLM available — fill in this week's content manually."


def test_docx_rollover_with_a_shorter_course_drops_the_trailing_rows(tmp_path):
    source = build_syllabus_with_full_schedule(tmp_path / "syllabus.docx")
    course_data = docx_extractor.extract_course_data(source)
    new_course_data, _ = roll_over_course(
        course_data, target_term="Fall 2027", calendars=CALENDARS, target_duration_weeks=3
    )
    output = tmp_path / "out" / "syllabus.docx"
    write_rolled_over_docx(source, new_course_data, output)

    reread = [w for w in docx_extractor.extract_course_data(output).weeks if not w.is_break]
    assert [w.week_number for w in reread] == [1, 2, 3]
