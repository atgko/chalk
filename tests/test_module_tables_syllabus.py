"""A longer real-world layout (16 weeks, schedule split into module
tables, month-name dates) — the demo's OSC 6660 sample
(chalk.demo_content.build_osc6660_module_syllabus)."""

import datetime as dt

from docx import Document

from chalk.demo_content import OSC6660_OBJECTIVES, build_osc6660_module_syllabus
from chalk.extractors import docx_extractor
from chalk.rollover.docx_rollover import write_rolled_over_docx
from chalk.rollover.preview import roll_over_course
from tests.test_rollover_season_change import CALENDARS


def _extract(tmp_path):
    return docx_extractor.extract_course_data(
        build_osc6660_module_syllabus(tmp_path / "syllabus.docx")
    )


def _roll(tmp_path, target_term="Spring 2027"):
    source = build_osc6660_module_syllabus(tmp_path / "syllabus.docx")
    new_course_data, preview = roll_over_course(
        docx_extractor.extract_course_data(source), target_term=target_term, calendars=CALENDARS
    )
    output = tmp_path / "out" / "syllabus.docx"
    write_rolled_over_docx(source, new_course_data, output)
    return output, preview


def _week_label_paragraphs(docx_path) -> list[str]:
    schedule_tables = Document(str(docx_path)).tables[1:]
    return [row.cells[0].paragraphs[0].text for table in schedule_tables for row in table.rows[2:]]


def test_course_details_come_from_the_unlabeled_header(tmp_path):
    course = _extract(tmp_path).course

    assert course.title == "Project Management in Healthcare"
    assert course.number == "OSC 6660"
    assert course.term == "Spring 2026"
    assert course.credits == 3
    assert course.meeting_pattern == "Tuesday 6:00-9:00 pm CRCC 105"
    assert course.instructor == "Alex Example & Sam Example"


def test_schedule_split_across_module_tables_is_read_as_one(tmp_path):
    course_data = _extract(tmp_path)

    assert [w.week_number for w in course_data.weeks] == list(range(1, 17))
    assert course_data.course.duration_weeks == 16
    assert course_data.weeks[0].date == dt.date(2026, 1, 6)
    assert course_data.weeks[-1].date == dt.date(2026, 4, 21)


def test_week_columns_split_into_topics_and_assignments(tmp_path):
    weeks = {w.week_number: w for w in _extract(tmp_path).weeks}

    assert weeks[2].topics == ["Project Charter", "Read: Article 2"]
    assert weeks[2].assignments == ["Project Artifact #2"]
    assert weeks[14].topics[0] == "Final Part 1 (Session 1): Final Presentations"


def test_objectives_skip_the_lead_in_and_stop_at_the_next_all_caps_heading(tmp_path):
    assert _extract(tmp_path).learning_objectives == OSC6660_OBJECTIVES


def test_weight_table_with_another_header_and_a_total_row(tmp_path):
    assessments = _extract(tmp_path).assessments

    assert [(a.name, a.weight) for a in assessments] == [
        ("Assignments", 0.3),
        ("Discussion Posts", 0.1),
        ("Final", 0.3),
        ("Participation", 0.3),
    ]


def test_rollover_keeps_the_month_name_style_and_text_after_the_date(tmp_path):
    output, _ = _roll(tmp_path)

    labels = _week_label_paragraphs(output)
    assert labels[0] == "Week 1 (Jan. 11)"
    assert labels[13] == "Week 14 (Apr. 12) – Final Part 1 (Session 1): Final Presentations"


def test_rollover_leaves_module_title_and_header_rows_alone(tmp_path):
    output, _ = _roll(tmp_path)

    tables = Document(str(output)).tables[1:]
    assert tables[0].rows[0].cells[0].text == "MODULE 1: Project Initiation"
    assert all(table.rows[1].cells[2].text == "Assignment Due" for table in tables)
    assert sum(len(table.rows) for table in tables) == 16 + 2 * 5


def test_rolled_over_file_re_extracts_cleanly(tmp_path):
    output, _ = _roll(tmp_path)

    reextracted = docx_extractor.extract_course_data(output)
    assert reextracted.course.term == "Spring 2027"
    assert reextracted.weeks[0].date == dt.date(2027, 1, 11)
    assert reextracted.course.duration_weeks == 16


def test_numbered_week_that_was_a_break_is_flagged(tmp_path):
    _, preview = _roll(tmp_path, target_term="Fall 2027")

    week10 = next(c for c in preview.week_changes if c.week_number == 10)
    assert any("Spring Break" in flag and "Fall 2027" in flag for flag in week10.flags)
