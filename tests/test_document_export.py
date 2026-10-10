"""Tests for chalk.generation.document_export: every generated document
type exports to Word and PDF with the assignment creator's layout, a
header naming the document type, a title, and real tables."""

import io

import pdfplumber
import pytest
from docx import Document

from chalk.generation.document_export import EXPORT_FORMATS, export_document

BANNER = "> **AI-generated draft** — review and edit before use. Generated 2026-10-10 using gpt-4o."
QUIZ = f"{BANNER}\n\n## Questions\nQ1. What is a subnet?\n\n## Answer Key\nA1. A network slice.\n"
RUBRIC = (
    f"{BANNER}\n\n"
    "| Criterion | Exemplary | Beginning |\n"
    "|---|---|---|\n"
    "| Design | Clear **plan** (40 pts) | Missing (10 pts) |\n"
    "| Report | Polished (60 pts) | Unclear (15 pts) |\n"
    "\nTotal: 100 points\n"
)


def _docx(data: bytes) -> Document:
    return Document(io.BytesIO(data))


def _pdf_text(data: bytes) -> str:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def test_formats_and_file_details():
    assert EXPORT_FORMATS == ("Word (.docx)", "PDF (.pdf)", "Markdown (.md)")
    _, ext, mime = export_document(QUIZ, "Word (.docx)", content_type="quiz", title="Week 3 — Quiz")
    assert (ext, mime) == ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    _, ext, mime = export_document(QUIZ, "PDF (.pdf)", content_type="quiz", title="Week 3 — Quiz")
    assert (ext, mime) == ("pdf", "application/pdf")


def test_markdown_export_is_the_draft_unchanged():
    data, ext, mime = export_document(QUIZ, "Markdown (.md)", content_type="quiz", title="Week 3 — Quiz")
    assert (data.decode("utf-8"), ext, mime) == (QUIZ, "md", "text/markdown")


def test_word_quiz_gets_header_title_and_no_banner():
    data, _, _ = export_document(QUIZ, "Word (.docx)", content_type="quiz", title="Week 3 — Quiz")
    document = _docx(data)
    paragraphs = [p.text for p in document.paragraphs]

    assert document.sections[0].header.paragraphs[0].text == "COURSE QUIZ"
    assert paragraphs[0] == "Week 3 — Quiz"
    assert "QUESTIONS" in paragraphs and "ANSWER KEY" in paragraphs
    assert not any("AI-generated draft" in p for p in paragraphs)


def test_a_draft_with_its_own_title_keeps_it():
    text = "# Lab 3: VLANs\n\n## Overview\nBuild a VLAN.\n"
    data, _, _ = export_document(text, "Word (.docx)", content_type="assignment", title="Lab 3 — Assignment")
    document = _docx(data)
    assert document.paragraphs[0].text == "Lab 3: VLANs"
    assert document.sections[0].header.paragraphs[0].text == "COURSE ASSIGNMENT"


@pytest.mark.parametrize(
    ("content_type", "header"),
    [("discussion", "DISCUSSION PROMPTS"), ("rubric", "GRADING RUBRIC"), ("summary", "MODULE SUMMARY")],
)
def test_each_type_names_itself_in_the_header(content_type, header):
    data, _, _ = export_document("Body text.", "Word (.docx)", content_type=content_type, title="T")
    assert _docx(data).sections[0].header.paragraphs[0].text == header


def test_word_rubric_table_becomes_a_real_table():
    data, _, _ = export_document(RUBRIC, "Word (.docx)", content_type="rubric", title="Lab 3 — Rubric")
    document = _docx(data)

    (table,) = document.tables
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    assert rows == [
        ["Criterion", "Exemplary", "Beginning"],
        ["Design", "Clear plan (40 pts)", "Missing (10 pts)"],
        ["Report", "Polished (60 pts)", "Unclear (15 pts)"],
    ]
    assert all(run.bold for run in table.rows[0].cells[0].paragraphs[0].runs)
    assert "Total: 100 points" in [p.text for p in document.paragraphs]
    assert not any(p.text.startswith("|") for p in document.paragraphs)


def test_pdf_quiz_has_header_title_and_no_banner():
    data, _, _ = export_document(QUIZ, "PDF (.pdf)", content_type="quiz", title="Week 3 - Quiz")
    text = _pdf_text(data)

    assert "COURSE QUIZ" in text
    assert "Week 3 - Quiz" in text
    assert "Q1. What is a subnet?" in text
    assert "AI-generated draft" not in text


def test_pdf_rubric_table_cells_are_laid_out_without_pipes():
    data, _, _ = export_document(RUBRIC, "PDF (.pdf)", content_type="rubric", title="Lab 3 Rubric")
    text = _pdf_text(data)

    assert "GRADING RUBRIC" in text
    assert "Clear plan (40 pts)" in text
    assert "|" not in text and "---" not in text


def test_pdf_escapes_angle_brackets_and_ampersands():
    data, _, _ = export_document(
        "Use <subnet> & mask.\n\n| A & B | <c> |\n|---|---|\n| x | y |\n",
        "PDF (.pdf)",
        content_type="summary",
        title="Week 1",
    )
    text = _pdf_text(data)
    assert "Use <subnet> & mask." in text
    assert "A & B" in text
