"""Programmatic synthetic .docx builders for extractor tests.

Each function builds the smallest Word document that exercises one
parsing rule, rather than relying on real (and non-shareable) faculty
syllabi — see DECISIONS.md and PLAN.md Section 5 ("Test Strategy").
"""

from __future__ import annotations

from pathlib import Path

from docx import Document


def _add_front_matter(
    doc,
    *,
    course: str = "IS 6640 Networking and Servers",
    number: str = "IS 6640",
    section: str = "090",
    credits: int = 3,
    term: str = "Fall 2026",
    instructor: str = "Dave Norwood",
    meeting_pattern: str = "Online async",
) -> None:
    doc.add_paragraph(f"Course: {course}")
    doc.add_paragraph(f"Course Number: {number}")
    doc.add_paragraph(f"Section: {section}")
    doc.add_paragraph(f"Credits: {credits}")
    doc.add_paragraph(f"Term: {term}")
    doc.add_paragraph(f"Instructor: {instructor}")
    doc.add_paragraph(f"Meeting Pattern: {meeting_pattern}")
    # A blank line, as real documents commonly have between sections —
    # exercises the extractor's blank-paragraph skip.
    doc.add_paragraph("")


def _add_learning_objectives(doc, *, heading: str = "Learning Objectives", objectives=None) -> None:
    objectives = objectives or ["Understand the fundamentals of IT network infrastructure"]
    doc.add_paragraph(heading, style="Heading 1")
    for objective in objectives:
        doc.add_paragraph(objective, style="List Bullet")


def _add_schedule_table(doc, rows: list[tuple[str, list[str]]]):
    """`rows` is a list of (week_label, second_cell_lines) tuples."""
    table = doc.add_table(rows=1, cols=2)
    header = table.rows[0].cells
    header[0].text = "Week"
    header[1].text = "Topics / Assignments"
    for label, lines in rows:
        row_cells = table.add_row().cells
        row_cells[0].text = label
        cell = row_cells[1]
        cell.text = lines[0] if lines else ""
        for line in lines[1:]:
            cell.add_paragraph(line)
    return table


def _add_assessments_table(doc, assessments: list[tuple[str, float]]):
    table = doc.add_table(rows=1, cols=2)
    header = table.rows[0].cells
    header[0].text = "Assessment"
    header[1].text = "Weight"
    for name, weight_pct in assessments:
        row_cells = table.add_row().cells
        row_cells[0].text = name
        row_cells[1].text = f"{weight_pct}%"
    return table


def _add_university_dates_table(doc, dates: list[tuple[str, str]]):
    table = doc.add_table(rows=1, cols=2)
    header = table.rows[0].cells
    header[0].text = "Event"
    header[1].text = "Date"
    for event, date_text in dates:
        row_cells = table.add_row().cells
        row_cells[0].text = event
        row_cells[1].text = date_text
    return table


def build_minimal_syllabus(path: Path) -> Path:
    """A complete minimal syllabus: front matter, learning objectives, a
    two-week schedule with one break week, an assessments table, and a
    university dates table. The happy-path fixture most extractor tests
    build on."""
    doc = Document()
    _add_front_matter(doc)
    _add_learning_objectives(doc)
    _add_schedule_table(
        doc,
        rows=[
            (
                "Week 1 (8/24)",
                [
                    "Course Introduction",
                    "Network+ Mod 1",
                    "Network+ Mod 2",
                    "Network+ Lab A Due",
                    "Note: Video Quizzes are Taken After Reading Each Chapter",
                ],
            ),
            ("Fall Break (10/10 - 10/18)", []),
            ("Week 2 (10/19)", ["Network+ Mod 3", "Network+ Lab B Due"]),
        ],
    )
    _add_assessments_table(doc, assessments=[("Video Quizzes", 5), ("Labs", 25)])
    _add_university_dates_table(
        doc,
        dates=[
            ("Classes begin", "8/24/2026"),
            ("Fall Break", "10/10 - 10/18"),
            ("Classes end", "12/10/2026"),
        ],
    )
    doc.save(str(path))
    return path


def build_syllabus_with_heading_after_objectives(path: Path) -> Path:
    """A heading (e.g. "Grading") right after the objectives bullet list,
    so the extractor's objectives collector has something to stop at."""
    doc = Document()
    _add_front_matter(doc)
    _add_learning_objectives(doc)
    doc.add_paragraph("Grading", style="Heading 1")
    doc.add_paragraph("This paragraph belongs to a later section, not the objectives.")
    _add_schedule_table(doc, rows=[("Week 1 (8/24)", ["Course Introduction"])])
    doc.save(str(path))
    return path


def build_syllabus_with_invalid_credits(path: Path) -> Path:
    doc = Document()
    _add_front_matter(doc, credits="three")  # type: ignore[arg-type]
    _add_schedule_table(doc, rows=[("Week 1 (8/24)", ["Course Introduction"])])
    doc.save(str(path))
    return path


def build_syllabus_with_unparseable_term_year(path: Path) -> Path:
    doc = Document()
    _add_front_matter(doc, term="Fall Semester")
    _add_schedule_table(doc, rows=[("Week 1 (8/24)", ["Course Introduction"])])
    doc.save(str(path))
    return path


def build_syllabus_with_unparseable_assessment_weight(path: Path) -> Path:
    doc = Document()
    _add_front_matter(doc)
    _add_schedule_table(doc, rows=[("Week 1 (8/24)", ["Course Introduction"])])
    _add_assessments_table(doc, assessments=[("Final Project", "TBD")])  # type: ignore[list-item]
    doc.save(str(path))
    return path


def build_syllabus_with_alternate_objectives_heading(path: Path) -> Path:
    doc = Document()
    _add_front_matter(doc)
    _add_learning_objectives(doc, heading="Course Outcome and Objectives")
    _add_schedule_table(doc, rows=[("Week 1 (8/24)", ["Course Introduction"])])
    doc.save(str(path))
    return path


def build_syllabus_without_university_dates_table(path: Path) -> Path:
    doc = Document()
    _add_front_matter(doc)
    _add_learning_objectives(doc)
    _add_schedule_table(doc, rows=[("Week 1 (8/24)", ["Course Introduction"])])
    _add_assessments_table(doc, assessments=[("Labs", 100)])
    doc.save(str(path))
    return path


def build_syllabus_with_no_schedule_table(path: Path) -> Path:
    doc = Document()
    _add_front_matter(doc)
    doc.add_paragraph("This document has no schedule table at all.")
    doc.save(str(path))
    return path


def build_syllabus_with_multiple_candidate_tables(path: Path) -> Path:
    """Two tables that both look like a schedule table (first data row
    matches the "Week N" pattern in each), forcing genuine ambiguity."""
    doc = Document()
    _add_front_matter(doc)
    _add_schedule_table(doc, rows=[("Week 1 (8/24)", ["Course Introduction"])])
    _add_schedule_table(doc, rows=[("Week 1 (8/24)", ["Duplicate table for ambiguity test"])])
    doc.save(str(path))
    return path


def build_syllabus_with_unresolvable_break_row(path: Path) -> Path:
    """A break row with no date range anywhere nearby — the extractor
    cannot recover date_start/date_end and must fail loudly rather than
    guess."""
    doc = Document()
    _add_front_matter(doc)
    _add_schedule_table(
        doc,
        rows=[
            ("Week 1 (8/24)", ["Course Introduction"]),
            ("Fall Break", []),  # no date range in label or body
        ],
    )
    doc.save(str(path))
    return path


def build_syllabus_missing_front_matter_field(path: Path, *, omit: str) -> Path:
    """A syllabus missing one required front-matter label (e.g. "Term").
    `omit` is one of: course, number, section, credits, term, instructor,
    meeting_pattern."""
    labels = {
        "course": "Course: IS 6640 Networking and Servers",
        "number": "Course Number: IS 6640",
        "section": "Section: 090",
        "credits": "Credits: 3",
        "term": "Term: Fall 2026",
        "instructor": "Instructor: Dave Norwood",
        "meeting_pattern": "Meeting Pattern: Online async",
    }
    doc = Document()
    for key, line in labels.items():
        if key != omit:
            doc.add_paragraph(line)
    _add_schedule_table(doc, rows=[("Week 1 (8/24)", ["Course Introduction"])])
    doc.save(str(path))
    return path


def build_syllabus_with_term_mismatch(path: Path) -> Path:
    """Reproduces the named sponsor bug: a Spring term label paired with
    an August (fall-pattern) Week 1 date."""
    doc = Document()
    _add_front_matter(doc, term="Spring 2026")
    _add_schedule_table(doc, rows=[("Week 1 (8/18)", ["Course Introduction"])])
    doc.save(str(path))
    return path


def build_syllabus_with_bold_week_label(path: Path) -> Path:
    """A schedule table whose Week 1 label cell has an explicitly bold
    run, to verify rollover preserves formatting rather than replacing it
    with a plain default run (python-docx#290)."""
    doc = Document()
    _add_front_matter(doc)
    table = doc.add_table(rows=1, cols=2)
    header = table.rows[0].cells
    header[0].text = "Week"
    header[1].text = "Topics / Assignments"
    row_cells = table.add_row().cells
    run = row_cells[0].paragraphs[0].add_run("Week 1 (8/24)")
    run.bold = True
    row_cells[1].text = "Course Introduction"
    doc.save(str(path))
    return path


def build_syllabus_with_multi_run_week_label(path: Path) -> Path:
    """A Week 1 label cell split across two runs with different
    formatting, to verify rollover clears the second run's text rather
    than leaving stale old-date text dangling alongside the new label."""
    doc = Document()
    _add_front_matter(doc)
    table = doc.add_table(rows=1, cols=2)
    header = table.rows[0].cells
    header[0].text = "Week"
    header[1].text = "Topics / Assignments"
    row_cells = table.add_row().cells
    paragraph = row_cells[0].paragraphs[0]
    first_run = paragraph.add_run("Week 1 ")
    first_run.bold = True
    second_run = paragraph.add_run("(8/24)")
    second_run.italic = True
    row_cells[1].text = "Course Introduction"
    doc.save(str(path))
    return path


def build_password_protected_docx(path: Path) -> Path:
    """Not a real password-protected OOXML file (that requires proper
    CFBF encryption machinery) — writes deliberately corrupt bytes that
    python-docx cannot open as a package. That's the actual condition the
    extractor needs to handle: *some* file it can't open, whatever the
    precise cause (DECISIONS.md: caught via exception, no zip pre-check)."""
    path.write_bytes(b"not a real docx file, deliberately corrupt")
    return path
