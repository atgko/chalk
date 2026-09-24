"""Tests for chalk.rollover.docx_rollover — validates in-place cell
mutation (formatting preservation), preserved topics/assignments, and
archive-before-write behavior.

new_course_data here is built by hand from the extracted source data
rather than via roll_over_course(), to isolate "does the writer apply a
given CourseData correctly" from "does roll_over_course compute the
right dates" (covered separately in test_rollover_preview.py).
"""

import datetime as dt

from docx import Document

from chalk.extractors.docx_extractor import extract_course_data
from chalk.metrics import read_events_by_type
from chalk.rollover.docx_rollover import write_rolled_over_docx
from tests.fixtures.docx_builder import (
    build_minimal_syllabus,
    build_syllabus_with_bold_week_label,
    build_syllabus_with_multi_run_week_label,
)


def _rolled_over_copy(course_data, *, week_dates: dict, break_dates=None, term: str = "Fall 2027"):
    new_weeks = []
    for week in course_data.weeks:
        if week.is_break:
            if break_dates is not None:
                new_start, new_end = break_dates
                new_weeks.append(
                    week.model_copy(
                        update={
                            "date_start": new_start,
                            "date_end": new_end,
                            "label": (
                                f"Fall Break ({new_start.month}/{new_start.day} - "
                                f"{new_end.month}/{new_end.day})"
                            ),
                        }
                    )
                )
            else:
                new_weeks.append(week)
            continue

        new_date = week_dates.get(week.week_number)
        if new_date is None:
            new_weeks.append(week)
            continue
        new_label = f"Week {week.week_number} ({new_date.month}/{new_date.day})"
        new_weeks.append(week.model_copy(update={"date": new_date, "label": new_label}))

    new_course_info = course_data.course.model_copy(update={"term": term})
    return course_data.model_copy(update={"course": new_course_info, "weeks": new_weeks})


def test_updates_week_label_dates_in_place(tmp_path):
    source_path = build_minimal_syllabus(tmp_path / "syllabus.docx")
    course_data = extract_course_data(source_path)

    new_course_data = _rolled_over_copy(
        course_data,
        week_dates={1: dt.date(2027, 8, 23), 2: dt.date(2027, 10, 18)},
        break_dates=(dt.date(2027, 10, 9), dt.date(2027, 10, 17)),
    )

    output_path = tmp_path / "outputs" / "syllabus.docx"
    write_rolled_over_docx(source_path, new_course_data, output_path)

    rolled_data = extract_course_data(output_path)
    assert rolled_data.weeks[0].date == dt.date(2027, 8, 23)
    week2 = next(week for week in rolled_data.weeks if week.week_number == 2)
    assert week2.date == dt.date(2027, 10, 18)
    break_week = next(week for week in rolled_data.weeks if week.is_break)
    assert break_week.date_start == dt.date(2027, 10, 9)
    assert break_week.date_end == dt.date(2027, 10, 17)


def test_term_label_is_updated_so_reextraction_uses_the_right_year(tmp_path):
    source_path = build_minimal_syllabus(tmp_path / "syllabus.docx")
    course_data = extract_course_data(source_path)

    new_course_data = _rolled_over_copy(
        course_data, week_dates={1: dt.date(2027, 8, 23)}, term="Fall 2027"
    )

    output_path = tmp_path / "outputs" / "syllabus.docx"
    write_rolled_over_docx(source_path, new_course_data, output_path)

    rolled_data = extract_course_data(output_path)
    assert rolled_data.course.term == "Fall 2027"
    # The whole point: without updating the Term line, re-extraction would
    # derive 2026 from the stale term label instead of 2027 from here.
    assert rolled_data.weeks[0].date == dt.date(2027, 8, 23)


def test_topics_assignments_and_notes_are_preserved_untouched(tmp_path):
    source_path = build_minimal_syllabus(tmp_path / "syllabus.docx")
    course_data = extract_course_data(source_path)

    new_course_data = _rolled_over_copy(course_data, week_dates={1: dt.date(2027, 8, 23)})

    output_path = tmp_path / "outputs" / "syllabus.docx"
    write_rolled_over_docx(source_path, new_course_data, output_path)

    week1 = extract_course_data(output_path).weeks[0]
    assert week1.topics == ["Course Introduction", "Network+ Mod 1", "Network+ Mod 2"]
    assert week1.assignments == ["Network+ Lab A Due"]
    assert week1.notes == "Video Quizzes are Taken After Reading Each Chapter"


def test_university_dates_table_is_updated(tmp_path):
    source_path = build_minimal_syllabus(tmp_path / "syllabus.docx")
    course_data = extract_course_data(source_path)

    new_course_data = _rolled_over_copy(course_data, week_dates={1: dt.date(2027, 8, 23)})
    new_university_dates = [
        entry.model_copy(update={"date": dt.date(2027, 8, 23)})
        if entry.event == "Classes begin"
        else entry
        for entry in new_course_data.university_dates
    ]
    new_course_data = new_course_data.model_copy(update={"university_dates": new_university_dates})

    output_path = tmp_path / "outputs" / "syllabus.docx"
    write_rolled_over_docx(source_path, new_course_data, output_path)

    rolled_data = extract_course_data(output_path)
    classes_begin = next(d for d in rolled_data.university_dates if d.event == "Classes begin")
    assert classes_begin.date == dt.date(2027, 8, 23)


def test_bold_formatting_on_the_week_label_is_preserved(tmp_path):
    source_path = build_syllabus_with_bold_week_label(tmp_path / "syllabus.docx")
    course_data = extract_course_data(source_path)

    new_course_data = _rolled_over_copy(course_data, week_dates={1: dt.date(2027, 8, 23)})

    output_path = tmp_path / "outputs" / "syllabus.docx"
    write_rolled_over_docx(source_path, new_course_data, output_path)

    result_document = Document(str(output_path))
    label_cell = result_document.tables[0].rows[1].cells[0]
    run = label_cell.paragraphs[0].runs[0]
    assert run.text == "Week 1 (8/23)"
    assert run.bold is True


def test_multi_run_week_label_has_the_second_run_cleared(tmp_path):
    source_path = build_syllabus_with_multi_run_week_label(tmp_path / "syllabus.docx")
    course_data = extract_course_data(source_path)

    new_course_data = _rolled_over_copy(course_data, week_dates={1: dt.date(2027, 8, 23)})

    output_path = tmp_path / "outputs" / "syllabus.docx"
    write_rolled_over_docx(source_path, new_course_data, output_path)

    result_document = Document(str(output_path))
    label_cell = result_document.tables[0].rows[1].cells[0]
    runs = label_cell.paragraphs[0].runs
    assert runs[0].text == "Week 1 (8/23)"
    assert runs[0].bold is True
    assert runs[1].text == ""


def test_university_date_row_with_no_matching_event_is_left_alone(tmp_path):
    source_path = build_minimal_syllabus(tmp_path / "syllabus.docx")
    course_data = extract_course_data(source_path)

    new_course_data = _rolled_over_copy(course_data, week_dates={1: dt.date(2027, 8, 23)})
    # Drop "Classes begin" from the new data entirely -- the doc's row for
    # it has nothing to match against and must be left untouched, not
    # crash.
    new_course_data = new_course_data.model_copy(
        update={
            "university_dates": [
                entry for entry in new_course_data.university_dates if entry.event != "Classes begin"
            ]
        }
    )

    output_path = tmp_path / "outputs" / "syllabus.docx"
    write_rolled_over_docx(source_path, new_course_data, output_path)

    result_document = Document(str(output_path))
    dates_table = next(
        t for t in result_document.tables if t.rows[0].cells[0].text.strip() == "Event"
    )
    first_data_row = dates_table.rows[1]
    assert first_data_row.cells[0].text.strip() == "Classes begin"
    assert first_data_row.cells[1].text.strip() == "8/24/2026"  # untouched, stale but not crashed


def test_original_source_file_is_never_modified(tmp_path):
    source_path = build_minimal_syllabus(tmp_path / "syllabus.docx")
    original_bytes = source_path.read_bytes()
    course_data = extract_course_data(source_path)

    new_course_data = _rolled_over_copy(course_data, week_dates={1: dt.date(2027, 8, 23)})
    write_rolled_over_docx(source_path, new_course_data, tmp_path / "outputs" / "syllabus.docx")

    assert source_path.read_bytes() == original_bytes


def test_existing_output_is_archived_before_being_overwritten(tmp_path):
    source_path = build_minimal_syllabus(tmp_path / "syllabus.docx")
    course_data = extract_course_data(source_path)
    new_course_data = _rolled_over_copy(course_data, week_dates={1: dt.date(2027, 8, 23)})

    output_path = tmp_path / "outputs" / "syllabus.docx"
    output_path.parent.mkdir(parents=True)
    output_path.write_text("stale previous version", encoding="utf-8")

    write_rolled_over_docx(source_path, new_course_data, output_path)

    archive_dir = output_path.parent / ".archive"
    archived_files = list(archive_dir.iterdir())
    assert len(archived_files) == 1
    assert archived_files[0].read_text(encoding="utf-8") == "stale previous version"

    rolled_data = extract_course_data(output_path)
    assert rolled_data.weeks[0].date == dt.date(2027, 8, 23)


def test_archive_event_is_logged_when_eval_log_path_given(tmp_path):
    source_path = build_minimal_syllabus(tmp_path / "syllabus.docx")
    course_data = extract_course_data(source_path)
    new_course_data = _rolled_over_copy(course_data, week_dates={1: dt.date(2027, 8, 23)})

    output_path = tmp_path / "outputs" / "syllabus.docx"
    output_path.parent.mkdir(parents=True)
    output_path.write_bytes(b"stale")
    log_path = tmp_path / "eval-log.json"

    write_rolled_over_docx(source_path, new_course_data, output_path, eval_log_path=log_path)

    assert len(read_events_by_type(log_path, "archive")) == 1
