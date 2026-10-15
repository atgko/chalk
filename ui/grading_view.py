from chalk.grading.export import export_word, export_pdf
"""AI Grading Assistant interface."""

from io import BytesIO

import streamlit as st

from chalk.grading.workflow import analyze_by_criterion
from chalk.llm_client import LLMProviderError


def extract_submission(uploaded_file):
    """Extract text from PDF, Word, or TXT submissions."""
    data = uploaded_file.getvalue()
    filename = uploaded_file.name.lower()

    if filename.endswith(".txt"):
        return data.decode("utf-8-sig")

    if filename.endswith(".docx"):
        from docx import Document

        document = Document(BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs)

    if filename.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    raise ValueError("Unsupported file type.")


def render(paths):
    """Render the grading assistant workflow."""
    st.header("AI Grading Assistant")
    st.caption(
        "Upload student work and generate rubric-based grading "
        "recommendations for instructor review."
    )

    st.info(
        "This feature provides grading suggestions. "
        "Instructors review and approve all final grades."
    )

    uploaded_file = st.file_uploader(
        "Upload student submission",
        type=["pdf", "docx", "txt"],
        key="grading_submission",
    )

    rubric = st.text_area(
        "Assignment instructions and grading rubric",
        height=200,
        placeholder="Paste the assignment requirements and rubric here...",
        key="grading_rubric",
    )

    instructions = st.text_area(
        "Additional grading instructions (optional)",
        height=100,
        placeholder="Enter any specific grading expectations...",
        key="grading_instructions",
    )

    ready = uploaded_file is not None and bool(rubric.strip())

    if st.button(
        "Analyze Submission",
        disabled=not ready,
        type="primary",
    ):
        try:
            with st.spinner("Analyzing student submission..."):
                submission = extract_submission(uploaded_file)

                if not submission.strip():
                    st.error(
                        "No readable text was found in the submission. "
                        "Scanned PDFs may require OCR."
                    )
                    return

                feedback = analyze_by_criterion(
                    submission=submission,
                    rubric=rubric,
                    instructions=instructions,
                )

            st.session_state["grading_feedback"] = feedback

        except (ValueError, UnicodeError) as exc:
            st.error(str(exc))
        except LLMProviderError as exc:
            st.error(str(exc))
        except Exception:
            st.error(
                "The submission could not be processed. "
                "Check the file format and try again."
            )

    if st.session_state.get("grading_feedback"):
        result = st.session_state["grading_feedback"]

        st.divider()
        st.subheader("AI Grading Recommendations")
        st.warning(
            "These are preliminary AI recommendations. "
            "The instructor must review and approve all final grades."
        )

        earned = result["earned_total"]
        possible = result["possible_total"]
        percentage = result["percentage"]

        st.metric(
            "Suggested Overall Grade",
            f"{earned:g} / {possible:g}",
            f"{percentage:.1f}%",
            delta_color="off",
        )

        st.subheader("Rubric Breakdown")

        for criterion in result["criteria"]:
            name = criterion["name"]
            score = criterion["earned"]
            maximum = criterion["possible"]

            st.markdown(
                f"**{name}: {score:g} / {maximum:g} points**"
            )
            st.progress(min(1.0, score / maximum))
            st.write(criterion["reasoning"])

            deduction = maximum - score
            comment = criterion.get("deduction_comment", "")

            if deduction > 0:
                st.markdown(
                    f"**Suggested deduction: -{deduction:g} points**"
                )
                if comment:
                    st.info(comment)
                else:
                    st.warning(
                        "Deduction explanation unavailable. "
                        "Instructor review required."
                    )

        st.divider()
        st.subheader("Student Feedback")

        student_feedback = result["student_feedback"]
        st.text_area(
            "Copy this comment into Canvas",
            value=student_feedback,
            height=150,
            key="student_feedback_display",
        )

        st.divider()
        st.subheader("Instructor Review Notes")
        st.write(result["instructor_notes"] or "No additional notes provided.")

        report = (
            f"# AI Grading Recommendations\\n\\n"
            f"Suggested Grade: {earned:g}/{possible:g} "
            f"({percentage:.1f}%)\\n\\n"
            "## Rubric Breakdown\\n\\n"
        )

        for criterion in result["criteria"]:
            report += (
                f"### {criterion['name']}\\n"
                f"Score: {criterion['earned']:g}/"
                f"{criterion['possible']:g}\\n\\n"
                f"{criterion['reasoning']}\\n\\n"
                + (
                    f"Suggested deduction: "
                    f"{criterion['possible'] - criterion['earned']:g} points\\n\\n"
                    f"{criterion['deduction_comment']}\\n\\n"
                    if criterion['possible'] > criterion['earned']
                    else ""
                )
            )

        report += (
            f"## Student Feedback\\n\\n{student_feedback}\\n\\n"
            f"## Instructor Review Notes\\n\\n"
            f"{result['instructor_notes']}\\n\\n"
            "All grades are preliminary recommendations "
            "requiring instructor approval.\\n"
        )

        st.subheader("Download Grading Report")

        word_data = export_word(result)
        pdf_data = export_pdf(result)

        col1, col2, col3 = st.columns(3)

        with col1:
            st.download_button(
                "Download Word",
                data=word_data,
                file_name="grading_feedback.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

        with col2:
            st.download_button(
                "Download PDF",
                data=pdf_data,
                file_name="grading_feedback.pdf",
                mime="application/pdf",
            )

        with col3:
            st.download_button(
                "Download Markdown",
                data=report,
                file_name="grading_feedback.md",
                mime="text/markdown",
            )


