"""Export preliminary grading recommendations."""

from io import BytesIO


def export_word(result):
    """Generate a Word grading report."""
    from docx import Document

    doc = Document()
    doc.add_heading("AI Grading Recommendations", 0)

    doc.add_paragraph(
        "Preliminary AI recommendations. "
        "Instructor approval is required."
    )

    earned = result["earned_total"]
    possible = result["possible_total"]
    percentage = result["percentage"]

    doc.add_heading("Suggested Overall Grade", level=1)
    doc.add_paragraph(
        f"{earned:g} / {possible:g} ({percentage:.1f}%)"
    )

    doc.add_heading("Rubric Breakdown", level=1)

    for criterion in result["criteria"]:
        name = criterion["name"]
        score = criterion["earned"]
        maximum = criterion["possible"]

        doc.add_heading(
            f"{name}: {score:g}/{maximum:g}",
            level=2,
        )
        doc.add_paragraph(criterion["reasoning"])

        if score < maximum:
            deduction = maximum - score
            doc.add_paragraph(
                f"Suggested deduction: {deduction:g} points"
            )
            doc.add_paragraph(
                criterion.get("deduction_comment")
                or "Instructor review required."
            )

    doc.add_heading("Student Feedback", level=1)
    doc.add_paragraph(result["student_feedback"])

    doc.add_heading("Instructor Review Notes", level=1)
    doc.add_paragraph(
        result["instructor_notes"] or "No additional notes."
    )

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def export_pdf(result):
    """Generate a PDF grading report."""
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from xml.sax.saxutils import escape

    output = BytesIO()
    doc = SimpleDocTemplate(output)
    styles = getSampleStyleSheet()
    story = []

    def add_text(value, style="Normal"):
        story.append(Paragraph(escape(str(value)), styles[style]))
        story.append(Spacer(1, 8))

    add_text("AI Grading Recommendations", "Title")
    add_text(
        "Preliminary AI recommendations. "
        "Instructor approval is required."
    )

    earned = result["earned_total"]
    possible = result["possible_total"]
    percentage = result["percentage"]

    add_text("Suggested Overall Grade", "Heading1")
    add_text(f"{earned:g}/{possible:g} ({percentage:.1f}%)", "Heading2")
    add_text("Rubric Breakdown", "Heading1")

    for criterion in result["criteria"]:
        score = criterion["earned"]
        maximum = criterion["possible"]

        add_text(
            f"{criterion['name']}: {score:g}/{maximum:g}",
            "Heading2",
        )
        add_text(criterion["reasoning"])

        if score < maximum:
            deduction = maximum - score
            add_text(
                f"Suggested deduction: {deduction:g} points",
                "Heading3",
            )
            add_text(
                criterion.get("deduction_comment")
                or "Instructor review required."
            )

    add_text("Student Feedback", "Heading1")
    add_text(result["student_feedback"])

    add_text("Instructor Review Notes", "Heading1")
    add_text(
        result["instructor_notes"] or "No additional notes."
    )

    doc.build(story)
    return output.getvalue()
