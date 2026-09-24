"""Word (.docx) syllabus extraction (F-01, Word path).

Locates the schedule table by the "Week N (M/DD)" heuristic in its first
column, reads front-matter metadata from labeled "Label: value" lines
preceding the table, learning objectives from a bulleted list under one
of two known headings, assessments from a two-column weight table, and
university dates from an optional two-column dates table.

Never writes to the source document — it is only ever opened for reading
(PRD section 6.1: "Do not modify any part of the document during
extraction").

The front-matter and cell-classification heuristics below are an MVP
convention, not something the PRD specifies precisely — they haven't yet
been validated against a real syllabus (that happens with the Section 15
test corpus, before the Nov 8 readiness meeting). See PLAN.md's "Word
format variability" risk.
"""

from __future__ import annotations

import datetime as dt
import re
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from chalk.errors import AmbiguousTableError, ExtractionError, ProtectedFileError
from chalk.extractors._dates import DATE_RANGE_RE, FULL_DATE_RE, year_from_term
from chalk.extractors._front_matter import finalize_front_matter, parse_labeled_line
from chalk.models import Assessment, CourseData, CourseInfo, UniversityDate, Week

# "Week 1 (8/24)", "Week 12 (12/1)" — identifies a schedule-table row as a
# regular (non-break) week.
WEEK_LABEL_RE = re.compile(r"week\s*(\d+)\s*\((\d{1,2})/(\d{1,2})\)", re.IGNORECASE)

_WEIGHT_PERCENT_RE = re.compile(r"(\d+(?:\.\d+)?)\s*%")

_OBJECTIVES_HEADINGS = {"learning objectives", "course outcome and objectives"}

_DUE_RE = re.compile(r"\bdue\b", re.IGNORECASE)
_NOTE_PREFIX_RE = re.compile(r"^notes?:\s*", re.IGNORECASE)


def extract_course_data(docx_path) -> CourseData:
    """Parse a Word syllabus into a CourseData object.

    Raises ProtectedFileError if the file can't be opened at all (usually
    password protection), AmbiguousTableError if zero or multiple
    candidate schedule tables are found, or ExtractionError for any other
    recognized parsing failure. Never writes to docx_path.
    """
    try:
        document = Document(str(docx_path))
    except (PackageNotFoundError, BadZipFile) as exc:
        raise ProtectedFileError(
            "This Word file is protected. Remove the password in Word "
            "(Review > Protect Document) and re-upload."
        ) from exc

    front_matter = _extract_front_matter(document)
    year = _year_from_term_or_raise(front_matter["term"])

    schedule_table = _find_schedule_table(document)
    weeks = _extract_weeks(schedule_table, year)

    course_info = CourseInfo(
        title=front_matter["title"],
        number=front_matter["number"],
        section=front_matter["section"],
        credits=front_matter["credits"],
        term=front_matter["term"],
        term_start=_course_start_date(weeks),
        term_end=_course_end_date(weeks),
        duration_weeks=sum(1 for week in weeks if not week.is_break),
        meeting_pattern=front_matter["meeting_pattern"],
        instructor=front_matter["instructor"],
        source_format="word",
        source_file=str(docx_path),
        extracted_at=dt.datetime.now(dt.timezone.utc),
    )

    return CourseData(
        course=course_info,
        learning_objectives=_extract_learning_objectives(document),
        assessments=_extract_assessments(document),
        weeks=weeks,
        university_dates=_extract_university_dates(document, year),
    )


# ---- Schedule table location -------------------------------------------


def _find_schedule_table(document):
    candidates = [table for table in document.tables if looks_like_schedule_table(table)]

    if len(candidates) == 1:
        return candidates[0]

    if len(candidates) == 0:
        raise AmbiguousTableError(
            "No schedule table found. Make sure your syllabus has a table "
            "with 'Week' in the first column, then try again.",
            candidates=[],
        )

    raise AmbiguousTableError(
        "Multiple possible schedule tables were found. Choose the correct one below.",
        candidates=[_table_preview(table) for table in candidates],
    )


def looks_like_schedule_table(table) -> bool:
    # Skip row 0 (assumed header) — at least one data row's first cell
    # must match the "Week N (M/DD)" pattern.
    return any(WEEK_LABEL_RE.search(row.cells[0].text) for row in table.rows[1:])


def _table_preview(table) -> str:
    if not table.rows:  # pragma: no cover - looks_like_schedule_table already requires rows
        return ""
    return " | ".join(cell.text for cell in table.rows[0].cells)


# ---- Front matter --------------------------------------------------------


def _extract_front_matter(document) -> dict:
    raw_fields: dict[str, str] = {}
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        parsed = parse_labeled_line(text)
        if parsed:
            field_name, value = parsed
            if field_name not in raw_fields:
                raw_fields[field_name] = value

    return finalize_front_matter(raw_fields)


def _year_from_term_or_raise(term: str) -> int:
    year = year_from_term(term)
    if year is None:
        raise ExtractionError(f"Could not determine the academic year from the term '{term}'.")
    return year


# ---- Schedule weeks -------------------------------------------------------


def _extract_weeks(table, year: int) -> list[Week]:
    weeks: list[Week] = []
    for row in table.rows[1:]:  # skip header
        label_text = row.cells[0].text.strip()
        second_cell_lines = _cell_lines(row.cells[1]) if len(row.cells) > 1 else []
        match = WEEK_LABEL_RE.search(label_text)

        if match:
            weeks.append(_build_regular_week(match, label_text, second_cell_lines, year))
        else:
            weeks.append(_build_break_week(label_text, second_cell_lines, year))

    if not weeks:  # pragma: no cover - _find_schedule_table already guarantees >=1 data row
        raise ExtractionError(
            "The schedule table was found but contains no week rows. Make "
            "sure it has at least one row below the header."
        )
    return weeks


def _build_regular_week(match, label_text: str, second_cell_lines: list[str], year: int) -> Week:
    week_number = int(match.group(1))
    month, day = int(match.group(2)), int(match.group(3))
    topics, assignments, notes = _classify_lines(second_cell_lines)
    return Week(
        week_number=week_number,
        date=dt.date(year, month, day),
        label=label_text,
        is_break=False,
        topics=topics,
        assignments=assignments,
        notes=notes,
    )


def _build_break_week(label_text: str, second_cell_lines: list[str], year: int) -> Week:
    combined_text = label_text + " " + " ".join(second_cell_lines)
    range_match = DATE_RANGE_RE.search(combined_text)
    if not range_match:
        raise ExtractionError(
            f"Could not determine the dates for the break row '{label_text}'. "
            "Include a date range like '(10/10 - 10/18)' next to the label."
        )
    m1, d1, m2, d2 = (int(group) for group in range_match.groups())
    return Week(
        week_number=None,
        date_start=dt.date(year, m1, d1),
        date_end=dt.date(year, m2, d2),
        label=label_text,
        is_break=True,
        topics=[],
        assignments=[],
    )


def _cell_lines(cell) -> list[str]:
    return [text for p in cell.paragraphs if (text := p.text.strip())]


def _classify_lines(lines: list[str]) -> tuple[list[str], list[str], str | None]:
    """Split a schedule cell's lines into topics, assignments, and an
    optional notes string.

    MVP heuristic: a line explicitly prefixed "Note:"/"Notes:" becomes
    part of the week's notes; a line containing the word "due" is an
    assignment; everything else is a topic.
    """
    topics: list[str] = []
    assignments: list[str] = []
    notes_parts: list[str] = []

    for line in lines:
        note_match = _NOTE_PREFIX_RE.match(line)
        if note_match:
            notes_parts.append(line[note_match.end() :].strip())
        elif _DUE_RE.search(line):
            assignments.append(line)
        else:
            topics.append(line)

    notes = " ".join(notes_parts) if notes_parts else None
    return topics, assignments, notes


def _course_start_date(weeks: list[Week]) -> dt.date:
    first = weeks[0]
    return first.date if first.date is not None else first.date_start


def _course_end_date(weeks: list[Week]) -> dt.date:
    last = weeks[-1]
    return last.date if last.date is not None else last.date_end


# ---- Learning objectives --------------------------------------------------


def _extract_learning_objectives(document) -> list[str]:
    objectives: list[str] = []
    collecting = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if text.lower().rstrip(":") in _OBJECTIVES_HEADINGS:
            collecting = True
            continue
        if not collecting:
            continue
        if paragraph.style is not None and paragraph.style.name.startswith("Heading"):
            break
        objectives.append(text)
    return objectives


# ---- Assessments table -----------------------------------------------------


def _extract_assessments(document) -> list[Assessment]:
    for table in document.tables:
        if not table.rows:  # pragma: no cover - python-docx tables always have >=1 row
            continue
        header = [cell.text.strip().lower() for cell in table.rows[0].cells]
        if "assessment" in header and any("weight" in cell for cell in header):
            return _parse_assessment_rows(table)
    return []


def _parse_assessment_rows(table) -> list[Assessment]:
    assessments = []
    for row in table.rows[1:]:
        name = row.cells[0].text.strip()
        weight_match = _WEIGHT_PERCENT_RE.search(row.cells[1].text)
        if not weight_match:
            continue
        assessments.append(Assessment(name=name, weight=float(weight_match.group(1)) / 100))
    return assessments


# ---- University dates table ------------------------------------------------


def _extract_university_dates(document, year: int) -> list[UniversityDate]:
    for table in document.tables:
        if not table.rows:  # pragma: no cover - python-docx tables always have >=1 row
            continue
        header = [cell.text.strip().lower() for cell in table.rows[0].cells]
        if "event" in header and "date" in header:
            return _parse_university_date_rows(table, year)
    return []


def _parse_university_date_rows(table, year: int) -> list[UniversityDate]:
    dates = []
    for row in table.rows[1:]:
        event = row.cells[0].text.strip()
        date_text = row.cells[1].text.strip()

        range_match = DATE_RANGE_RE.search(date_text)
        if range_match:
            m1, d1, m2, d2 = (int(group) for group in range_match.groups())
            dates.append(
                UniversityDate(
                    event=event, date_start=dt.date(year, m1, d1), date_end=dt.date(year, m2, d2)
                )
            )
            continue

        full_match = FULL_DATE_RE.search(date_text)
        if full_match:
            month, day, full_year = (int(group) for group in full_match.groups())
            dates.append(UniversityDate(event=event, date=dt.date(full_year, month, day)))

    return dates
