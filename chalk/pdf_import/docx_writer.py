"""Write a PDF's blocks out as a Word document the existing Word path can
read: headings in a Heading style (the objectives reader stops at them),
bullets as List Bullet, and tables with one paragraph per cell line (the
schedule reader classifies a cell line by line).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from chalk.pdf_import.blocks import Block, Paragraph, Table

_STYLE_BY_KIND = {"heading": "Heading 2", "bullet": "List Bullet", "body": None}


def write_docx(blocks: list[Block], path: Path) -> Path:
    document = Document()
    for block in blocks:
        if isinstance(block, Paragraph):
            document.add_paragraph(block.text, style=_STYLE_BY_KIND[block.kind])
        else:
            _add_table(document, block)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def _add_table(document, table: Table) -> None:
    width = max(len(row) for row in table.rows)
    docx_table = document.add_table(rows=len(table.rows), cols=width)
    docx_table.style = "Table Grid"
    for docx_row, row in zip(docx_table.rows, table.rows):
        for cell, lines in zip(docx_row.cells, row):
            if not lines:
                continue
            cell.paragraphs[0].text = lines[0]
            for line in lines[1:]:
                cell.add_paragraph(line)
