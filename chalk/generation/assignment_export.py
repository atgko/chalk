"""Export student-facing assignments to Markdown, Word, or PDF."""

from __future__ import annotations

import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)


def _clean_line(line: str) -> str:
    """Remove basic Markdown formatting for document exports."""
    line = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
    line = re.sub(r"\*(.*?)\*", r"\1", line)
    return line.strip()


def _assignment_title(text: str) -> str:
    """Return the first level-one Markdown heading."""
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if line.startswith("# ") and not line.startswith("## "):
            return _clean_line(line[2:])
    return "Assignment"


def export_markdown(text: str) -> bytes:
    return text.encode("utf-8")


def _add_docx_header(document: Document) -> None:
    """Add a clean university-style header."""
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run("COURSE ASSIGNMENT")
    run.bold = True
    run.font.size = Pt(9)

    paragraph.paragraph_format.space_after = Pt(3)


def _style_docx(document: Document) -> None:
    """Apply readable academic-document typography."""
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size in (
        ("Title", 20),
        ("Heading 1", 14),
        ("Heading 2", 12),
        ("Heading 3", 11),
    ):
        style = styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True

    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 2"].paragraph_format.space_before = Pt(11)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)


def export_docx(text: str) -> bytes:
    """Create a polished student-facing Word assignment."""
    buffer = BytesIO()
    document = Document()

    _style_docx(document)
    _add_docx_header(document)

    title_written = False

    for raw_line in text.splitlines():
        line = raw_line.strip()

        if not line:
            continue

        if line.startswith("# ") and not line.startswith("## "):
            title = _clean_line(line[2:])
            paragraph = document.add_paragraph()
            paragraph.style = document.styles["Title"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(18)

            run = paragraph.add_run(title)
            run.bold = True
            title_written = True

        elif line.startswith("### "):
            document.add_heading(_clean_line(line[4:]), level=2)

        elif line.startswith("## "):
            document.add_heading(_clean_line(line[3:]).upper(), level=1)

        elif line.startswith(("- ", "* ")):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(_clean_line(line[2:]))

        elif re.match(r"^\d+\.\s", line):
            content = re.sub(r"^\d+\.\s*", "", line)
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run(_clean_line(content))

        elif line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            run = paragraph.add_run(_clean_line(line[2:]))
            run.italic = True

        else:
            document.add_paragraph(_clean_line(line))

    if not title_written:
        paragraph = document.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.save(buffer)
    return buffer.getvalue()


def _pdf_styles():
    styles = getSampleStyleSheet()

    styles.add(
        ParagraphStyle(
            name="AssignmentTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=19,
            leading=23,
            alignment=TA_CENTER,
            spaceAfter=18,
        )
    )

    styles.add(
        ParagraphStyle(
            name="AssignmentSection",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            spaceBefore=12,
            spaceAfter=6,
        )
    )

    styles.add(
        ParagraphStyle(
            name="AssignmentSubsection",
            parent=styles["Heading3"],
            fontName="Helvetica-Bold",
            fontSize=10.5,
            leading=13,
            spaceBefore=9,
            spaceAfter=4,
        )
    )

    styles.add(
        ParagraphStyle(
            name="AssignmentBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=10.5,
            leading=15,
            spaceAfter=7,
        )
    )

    styles.add(
        ParagraphStyle(
            name="AssignmentHeader",
            parent=styles["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=8.5,
            leading=10,
            alignment=TA_CENTER,
            spaceAfter=10,
        )
    )

    return styles


def _pdf_header_footer(canvas, doc) -> None:
    canvas.saveState()

    width, height = letter

    canvas.setFont("Helvetica-Bold", 8)
    canvas.drawCentredString(width / 2, height - 0.38 * inch, "COURSE ASSIGNMENT")

    canvas.setFont("Helvetica", 8)
    canvas.drawCentredString(width / 2, 0.38 * inch, f"Page {doc.page}")

    canvas.restoreState()


def export_pdf(text: str) -> bytes:
    """Create a polished student-facing PDF assignment."""
    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.8 * inch,
        leftMargin=0.8 * inch,
        topMargin=0.72 * inch,
        bottomMargin=0.65 * inch,
        title=_assignment_title(text),
    )

    styles = _pdf_styles()
    story = []
    pending_bullets = []
    pending_numbers = []

    def flush_lists():
        nonlocal pending_bullets, pending_numbers

        if pending_bullets:
            story.append(
                ListFlowable(
                    [
                        ListItem(
                            Paragraph(item, styles["AssignmentBody"]),
                            leftIndent=12,
                        )
                        for item in pending_bullets
                    ],
                    bulletType="bullet",
                    leftIndent=22,
                )
            )
            story.append(Spacer(1, 5))
            pending_bullets = []

        if pending_numbers:
            story.append(
                ListFlowable(
                    [
                        ListItem(
                            Paragraph(item, styles["AssignmentBody"]),
                            leftIndent=12,
                        )
                        for item in pending_numbers
                    ],
                    bulletType="1",
                    start="1",
                    leftIndent=22,
                )
            )
            story.append(Spacer(1, 5))
            pending_numbers = []

    for raw_line in text.splitlines():
        line = raw_line.strip()

        if not line:
            flush_lists()
            continue

        if line.startswith("- "):
            flush_lists()
            pending_bullets.append(_clean_line(line[2:]))
            continue

        if re.match(r"^\d+\.\s", line):
            flush_lists()
            pending_numbers.append(
                _clean_line(re.sub(r"^\d+\.\s*", "", line))
            )
            continue

        flush_lists()

        if line.startswith("# ") and not line.startswith("## "):
            story.append(
                Paragraph(
                    _clean_line(line[2:]),
                    styles["AssignmentTitle"],
                )
            )

        elif line.startswith("### "):
            story.append(
                Paragraph(
                    _clean_line(line[4:]),
                    styles["AssignmentSubsection"],
                )
            )

        elif line.startswith("## "):
            story.append(
                Paragraph(
                    _clean_line(line[3:]).upper(),
                    styles["AssignmentSection"],
                )
            )

        elif line.startswith("> "):
            story.append(
                Paragraph(
                    f"<i>{_clean_line(line[2:])}</i>",
                    styles["AssignmentBody"],
                )
            )

        else:
            story.append(
                Paragraph(
                    _clean_line(line),
                    styles["AssignmentBody"],
                )
            )

    flush_lists()

    document.build(
        story,
        onFirstPage=_pdf_header_footer,
        onLaterPages=_pdf_header_footer,
    )

    return buffer.getvalue()


def export_assignment(text: str, file_format: str) -> tuple[bytes, str, str]:
    """Export assignment text in the selected format."""
    if file_format == "Word (.docx)":
        return (
            export_docx(text),
            "docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    if file_format == "PDF (.pdf)":
        return export_pdf(text), "pdf", "application/pdf"

    return export_markdown(text), "md", "text/markdown"
