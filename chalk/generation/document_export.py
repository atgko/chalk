"""Export generated documents (assignments, quizzes, discussion prompts,
rubrics, summaries) to Word, PDF, or Markdown, all in the assignment
creator's layout: a small header naming the document type, a centered
title, upper-case section headings, and real tables for markdown tables.

Word and PDF are student-facing, so the AI-draft banner line is left out
of them (the saved markdown keeps it). A draft without its own `# Title`
gets the title it's given.
"""

from __future__ import annotations

import re
from collections.abc import Callable
from html import escape
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

EXPORT_FORMATS = ("Word (.docx)", "PDF (.pdf)", "Markdown (.md)")
HEADERS = {
    "assignment": "COURSE ASSIGNMENT",
    "quiz": "COURSE QUIZ",
    "discussion": "DISCUSSION PROMPTS",
    "rubric": "GRADING RUBRIC",
    "summary": "MODULE SUMMARY",
}
_DOCUMENT_NAMES = {
    "assignment": "Assignment",
    "quiz": "Quiz",
    "discussion": "Discussion Prompts",
    "rubric": "Rubric",
    "summary": "Module Summary",
}
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_BANNER_RE = re.compile(r"^>\s*\*\*AI-generated draft\*\*")
_NUMBERED_RE = re.compile(r"^\d+\.\s")
_TABLE_SEPARATOR_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?$")


def export_document(text: str, file_format: str, *, content_type: str, title: str) -> tuple[bytes, str, str]:
    """The draft as (file bytes, extension, MIME type) in the chosen format."""
    if file_format == "Markdown (.md)":
        return text.encode("utf-8"), "md", "text/markdown"
    header = HEADERS.get(content_type, "COURSE DOCUMENT")
    blocks = _parse(_prepare(text, title))
    if file_format == "PDF (.pdf)":
        return _export_pdf(blocks, header), "pdf", "application/pdf"
    return _export_docx(blocks, header), "docx", _DOCX_MIME


def can_export(content_type: str) -> bool:
    """Slides are Pandoc markdown for a slide deck, not a document."""
    return content_type in HEADERS


def document_title(content_type: str, subject: str) -> str:
    """The title a draft without its own `# Title` gets, e.g. "Week 3 — Quiz"."""
    return f"{subject} — {_DOCUMENT_NAMES.get(content_type, 'Document')}"


# ---- Markdown -> blocks ---------------------------------------------------------------


def _prepare(text: str, title: str) -> str:
    lines = [line for line in text.splitlines() if not _BANNER_RE.match(line.strip())]
    has_title = any(_is_title(line.strip()) for line in lines)
    if title and not has_title:
        lines = [f"# {title}", "", *lines]
    return "\n".join(lines)


def _is_title(line: str) -> bool:
    return line.startswith("# ")


def _parse(text: str) -> list[tuple[str, object]]:
    """(kind, content) blocks: title, section, subsection, bullet, number,
    quote, paragraph (a string each), and table (a list of rows)."""
    blocks: list[tuple[str, object]] = []
    table: list[list[str]] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if line.startswith("|"):
            if not _TABLE_SEPARATOR_RE.match(line):
                table.append([_clean(cell) for cell in line.strip("|").split("|")])
            continue
        if table:
            blocks.append(("table", table))
            table = []
        if line:
            blocks.append(_line_block(line))
        else:
            blocks.append(("blank", ""))
    if table:
        blocks.append(("table", table))
    return blocks


def _line_block(line: str) -> tuple[str, str]:
    if _is_title(line):
        return "title", _clean(line[2:])
    if line.startswith("### "):
        return "subsection", _clean(line[4:])
    if line.startswith("## "):
        return "section", _clean(line[3:]).upper()
    if line.startswith(("- ", "* ")):
        return "bullet", _clean(line[2:])
    if _NUMBERED_RE.match(line):
        return "number", _clean(_NUMBERED_RE.sub("", line))
    if line.startswith("> "):
        return "quote", _clean(line[2:])
    return "paragraph", _clean(line)


def _clean(text: str) -> str:
    """Drop basic markdown emphasis for document exports."""
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    return text.strip()


def _title_of(blocks: list[tuple[str, object]]) -> str:
    return next((str(content) for kind, content in blocks if kind == "title"), "Document")


# ---- Word -----------------------------------------------------------------------------


def _export_docx(blocks: list[tuple[str, object]], header: str) -> bytes:
    document = Document()
    _style_docx(document)
    _add_docx_header(document, header)
    for kind, content in blocks:
        _DOCX_WRITERS.get(kind, _docx_skip)(document, content)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _style_docx(document) -> None:
    """Readable academic-document typography."""
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        style = styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 2"].paragraph_format.space_before = Pt(11)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)


def _add_docx_header(document, header: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    paragraph = section.header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(header)
    run.bold = True
    run.font.size = Pt(9)
    paragraph.paragraph_format.space_after = Pt(3)


def _docx_title(document, text: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(18)
    paragraph.add_run(text).bold = True


def _docx_quote(document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.add_run(text).italic = True


def _docx_table(document, rows: list[list[str]]) -> None:
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width, style="Table Grid")
    for r, row in enumerate(rows):
        for c in range(width):
            run = table.cell(r, c).paragraphs[0].add_run(row[c] if c < len(row) else "")
            run.bold = r == 0
            run.font.size = Pt(9.5)
    document.add_paragraph()


def _docx_skip(document, content) -> None:
    return None


_DOCX_WRITERS: dict[str, Callable] = {
    "title": _docx_title,
    "section": lambda d, text: d.add_heading(text, level=1),
    "subsection": lambda d, text: d.add_heading(text, level=2),
    "bullet": lambda d, text: d.add_paragraph(text, style="List Bullet"),
    "number": lambda d, text: d.add_paragraph(text, style="List Number"),
    "quote": _docx_quote,
    "paragraph": lambda d, text: d.add_paragraph(text),
    "table": _docx_table,
}


# ---- PDF ------------------------------------------------------------------------------


def _export_pdf(blocks: list[tuple[str, object]], header: str) -> bytes:
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.8 * inch,
        leftMargin=0.8 * inch,
        topMargin=0.72 * inch,
        bottomMargin=0.65 * inch,
        title=_title_of(blocks),
    )
    styles = _pdf_styles()
    story = _pdf_story(blocks, styles, document.width)

    def decorate(canvas, doc) -> None:
        width, height = letter
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 8)
        canvas.drawCentredString(width / 2, height - 0.38 * inch, header)
        canvas.setFont("Helvetica", 8)
        canvas.drawCentredString(width / 2, 0.38 * inch, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buffer.getvalue()


def _pdf_story(blocks: list[tuple[str, object]], styles, width: float) -> list:
    story: list = []
    pending: list[str] = []
    pending_kind = ""

    def flush() -> None:
        nonlocal pending, pending_kind
        if pending:
            story.append(_pdf_list(pending, pending_kind, styles))
            story.append(Spacer(1, 5))
        pending, pending_kind = [], ""

    for kind, content in blocks:
        if kind in ("bullet", "number"):
            if kind != pending_kind:
                flush()
            pending.append(str(content))
            pending_kind = kind
            continue
        flush()
        if kind == "table":
            story += [_pdf_table(content, styles, width), Spacer(1, 8)]
        elif kind in _PDF_STYLE_FOR:
            text = escape(str(content), quote=False)
            markup = f"<i>{text}</i>" if kind == "quote" else text
            story.append(Paragraph(markup, styles[_PDF_STYLE_FOR[kind]]))
    flush()
    return story


_PDF_STYLE_FOR = {
    "title": "DocTitle",
    "section": "DocSection",
    "subsection": "DocSubsection",
    "quote": "DocBody",
    "paragraph": "DocBody",
}


def _pdf_list(items: list[str], kind: str, styles) -> ListFlowable:
    entries = [
        ListItem(Paragraph(escape(item, quote=False), styles["DocBody"]), leftIndent=12) for item in items
    ]
    if kind == "number":
        return ListFlowable(entries, bulletType="1", start="1", leftIndent=22)
    return ListFlowable(entries, bulletType="bullet", leftIndent=22)


def _pdf_table(rows: list[list[str]], styles, width: float) -> Table:
    columns = max(len(row) for row in rows)
    padded = [row + [""] * (columns - len(row)) for row in rows]
    cells = [
        [Paragraph(escape(text, quote=False), styles["DocTableHead" if r == 0 else "DocTable"]) for text in row]
        for r, row in enumerate(padded)
    ]
    table = Table(cells, colWidths=[width / columns] * columns, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    return table


def _pdf_styles():
    styles = getSampleStyleSheet()
    for name, parent, font, size, leading, extra in (
        ("DocTitle", "Title", "Helvetica-Bold", 19, 23, {"alignment": TA_CENTER, "spaceAfter": 18}),
        ("DocSection", "Heading2", "Helvetica-Bold", 11, 14, {"spaceBefore": 12, "spaceAfter": 6}),
        ("DocSubsection", "Heading3", "Helvetica-Bold", 10.5, 13, {"spaceBefore": 9, "spaceAfter": 4}),
        ("DocBody", "BodyText", "Helvetica", 10.5, 15, {"spaceAfter": 7}),
        ("DocTable", "BodyText", "Helvetica", 9, 12, {}),
        ("DocTableHead", "BodyText", "Helvetica-Bold", 9, 12, {}),
    ):
        styles.add(
            ParagraphStyle(
                name=name, parent=styles[parent], fontName=font, fontSize=size, leading=leading, **extra
            )
        )
    return styles
