"""Source-material context for generation prompts (DECISIONS.md): the
selected source/ files' text, concatenated, truncated to a token budget.
No embeddings or retrieval for the MVP.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from chalk.costs import CHARS_PER_TOKEN
from chalk.errors import GenerationError
from chalk.project import ProjectPaths, list_source_files

DEFAULT_TOKEN_BUDGET = 6000

NO_SOURCES_TEXT = (
    "No source materials were provided. Rely on the week's topics and the "
    "course learning objectives."
)


def read_source_text(path: Path) -> str:
    """Plain text from one source file. Returns "" for a file with no
    extractable text (e.g. a scanned PDF) rather than failing the whole
    generation."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages).strip()
        except PdfReadError:
            return ""
    if suffix == ".docx":
        document = Document(str(path))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            lines += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
        return "\n".join(line for line in lines if line.strip())
    return path.read_text(encoding="utf-8", errors="replace").strip()


def build_source_context(
    paths: ProjectPaths, filenames: list[str], *, token_budget: int = DEFAULT_TOKEN_BUDGET
) -> str:
    """The prompt's {source_material_context} block for `filenames`
    (names within source/). Unknown names raise GenerationError — they're
    never resolved as arbitrary paths."""
    if not filenames:
        return NO_SOURCES_TEXT

    available = set(list_source_files(paths))
    unknown = [name for name in filenames if name not in available]
    if unknown:
        raise GenerationError(
            f"These source files aren't in this project's source/ folder: {', '.join(unknown)}. "
            "Add them first."
        )

    sections = []
    for name in filenames:
        text = read_source_text(paths.source_dir / name) or "(No extractable text in this file.)"
        sections.append(f"--- Source: {name} ---\n{text}")
    context = "Source materials supplied by the instructor:\n\n" + "\n\n".join(sections)
    return _truncate(context, token_budget * CHARS_PER_TOKEN)


def _truncate(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    omitted = len(text) - max_chars
    return text[:max_chars] + f"\n\n[Source material truncated: {omitted:,} more characters omitted.]"
