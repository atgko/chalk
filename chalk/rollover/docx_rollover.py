"""Word (.docx) rollover writer (F-02, Word path).

Applies a rolled-over CourseData's new dates onto a fresh copy of the
original Word document — mutating each week-label cell's existing run
text in place rather than clearing and rebuilding the paragraph, so
formatting (bold, color, font) survives (python-openxml/python-docx#290;
see PLAN.md Risk #1). All other cell content (topics, assignments, notes)
is left untouched, per PRD section 6.2: "All other cell content
preserved." Quiz/exam/lab numbering embedded in that untouched text is
therefore never renumbered — flagging that for human review is
roll_over_course()'s job (chalk.rollover.preview), not this writer's.

chalk.rollover.preview.roll_over_course() must be called first; this
module only ever applies an already-computed CourseData to a document.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from chalk.archiving import archive_before_write
from chalk.extractors._front_matter import parse_labeled_line
from chalk.extractors.docx_extractor import WEEK_LABEL_RE, looks_like_schedule_table
from chalk.models import CourseData, UniversityDate
from chalk.rollover.preview import break_name


def write_rolled_over_docx(
    source_docx_path,
    new_course_data: CourseData,
    output_path,
    *,
    eval_log_path=None,
) -> None:
    """Write the rolled-over schedule into a fresh copy of the original
    Word document at `output_path`, archiving whatever was already there
    first (PRD's "Confirm rollover" step never writes without archiving).
    """
    document = Document(str(source_docx_path))

    _update_front_matter_term(document, new_course_data.course.term)

    new_dates_by_week_number = {
        week.week_number: week.date
        for week in new_course_data.weeks
        if not week.is_break and week.date is not None
    }
    new_break_labels_by_name = {
        break_name(week.label): week.label for week in new_course_data.weeks if week.is_break
    }

    for table in document.tables:
        if looks_like_schedule_table(table):
            _update_schedule_table(table, new_dates_by_week_number, new_break_labels_by_name)
        elif _looks_like_university_dates_table(table):
            _update_university_dates_table(table, new_course_data.university_dates)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    archive_before_write(output_path, eval_log_path=eval_log_path)
    document.save(str(output_path))


def _update_schedule_table(
    table, new_dates_by_week_number: dict, new_break_labels_by_name: dict
) -> None:
    for row in table.rows[1:]:  # skip header
        label_cell = row.cells[0]
        label_text = label_cell.text.strip()
        match = WEEK_LABEL_RE.search(label_text)

        if match:
            week_number = int(match.group(1))
            new_date = new_dates_by_week_number.get(week_number)
            if new_date is not None:
                new_label = f"Week {week_number} ({new_date.month}/{new_date.day})"
                _set_first_paragraph_text_preserving_format(label_cell, new_label)
        else:
            new_label = new_break_labels_by_name.get(break_name(label_text))
            if new_label is not None:
                _set_first_paragraph_text_preserving_format(label_cell, new_label)


def _update_front_matter_term(document, new_term: str) -> None:
    """Update the "Term: ..." front-matter line to the target term.

    Without this, re-extracting the written-out document would still see
    the old term label and derive the wrong calendar year from it
    (chalk.extractors._front_matter / docx_extractor's year_from_term),
    even though every week's date text was correctly updated. This is
    also the change PRD section 6.2's rollover preview shows first:
    "Term label: 'Fall 2026' -> 'Fall 2027'".
    """
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:  # pragma: no cover - none of our fixtures put a blank line before Term
            continue
        parsed = parse_labeled_line(text)
        if parsed and parsed[0] == "term":
            _set_paragraph_text_preserving_format(paragraph, f"Term: {new_term}")
            return


def _looks_like_university_dates_table(table) -> bool:
    if not table.rows:  # pragma: no cover - python-docx tables always have >=1 row
        return False
    header = [cell.text.strip().lower() for cell in table.rows[0].cells]
    return "event" in header and "date" in header


def _update_university_dates_table(table, university_dates: list[UniversityDate]) -> None:
    dates_by_event = {entry.event: entry for entry in university_dates}
    for row in table.rows[1:]:
        entry = dates_by_event.get(row.cells[0].text.strip())
        if entry is None:
            continue
        _set_first_paragraph_text_preserving_format(row.cells[1], _format_university_date(entry))


def _format_university_date(entry: UniversityDate) -> str:
    if entry.date is not None:
        return f"{entry.date.month}/{entry.date.day}/{entry.date.year}"
    return (
        f"{entry.date_start.month}/{entry.date_start.day} - "
        f"{entry.date_end.month}/{entry.date_end.day}"
    )


def _set_first_paragraph_text_preserving_format(cell, new_text: str) -> None:
    _set_paragraph_text_preserving_format(cell.paragraphs[0], new_text)


def _set_paragraph_text_preserving_format(paragraph, new_text: str) -> None:
    """Mutate a paragraph's existing runs in place, rather than clearing
    it and adding a fresh (unformatted) run. `cell.text = "..."` or a
    clear-and-rebuild would silently drop bold/color/font formatting —
    see python-openxml/python-docx#290."""
    runs = paragraph.runs
    if not runs:  # pragma: no cover - a paragraph with visible text always has >=1 run
        paragraph.add_run(new_text)
        return
    runs[0].text = new_text
    for extra_run in runs[1:]:
        extra_run.text = ""
